"""
Products, phases 1 and 2 — the markets a product serves, and changing the catalogue.

    python run_tests.py test_products_editing.py

What is proved (13 Oct 2026):
  · A product is created, edited, deactivated and deleted through the API. It
    is asked for a PRODUCT's required fields only — never the price matrix's
    or a pricing table's, which the same register module describes.
  · A SKU Code is one product's; a product can't end before it starts; a
    product a Lead names can't be deleted, and a number is never used twice.
  · Segment -> Segment Category -> Subsegment take several choices each, and
    a choice its parent does not offer is refused.
  · Only an admin or a developer changes the catalogue; anyone reads it.
  · Products import and export, a product's own fields only.

Needs the catalogue loaded and products_market_segments_metadata.py applied in
the database the test copy is made from.
"""

import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy import delete, func, select  # noqa: E402

import test_db  # noqa: E402

test_db.require_test_database()

from app import requirements  # noqa: E402
from app.auth import current_user  # noqa: E402
from app.clock import now_utc  # noqa: E402
from app.database import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402
from app.models import AuditLog, PicklistValue, PriceBook, PriceBookEdition, PriceBookEntry, Product, User, UserRole  # noqa: E402
from test_support import sign_in_as_admin  # noqa: E402

client = TestClient(app)
USER = sign_in_as_admin(app)
# test_support switches required fields off for the older scripts. A product's
# are what is being proved here; the Lead below is made with them off.
requirements.ENFORCED = True

failures: list[str] = []
created: dict[str, list[str]] = {"leads": [], "accounts": [], "products": []}

#: Someone with a role that is neither ADMIN nor DEVELOPER.
VIEWER = "USR-T-PR1"
#: Every SKU this script makes starts with this, so its cleanup can find them.
SKU = "TPROBE"


def check(label: str, condition: bool, detail: str = "") -> None:
    if condition:
        print(f"  pass  {label}")
    else:
        failures.append(f"{label}{' — ' + detail if detail else ''}")
        print(f"  FAIL  {label}{' — ' + detail if detail else ''}")


def code_of(response) -> str | None:
    try:
        detail = response.json().get("detail")
    except ValueError:
        return None
    return detail.get("code") if isinstance(detail, dict) else None


def missing_of(response) -> set[str]:
    try:
        detail = response.json().get("detail")
    except ValueError:
        return set()
    return {f["api_name"] for f in detail.get("fields", [])} if isinstance(detail, dict) else set()


def act_as_viewer() -> None:
    def _current_user():
        with SessionLocal() as db:
            user = db.get(User, VIEWER)
            _ = user.roles
            db.expunge(user)
            return user

    app.dependency_overrides[current_user] = _current_user


def choices(picklist: str) -> list[tuple[str, str]]:
    with SessionLocal() as db:
        rows = db.scalars(
            select(PicklistValue)
            .where(PicklistValue.picklist_key == picklist, PicklistValue.active.is_(True))
            .order_by(PicklistValue.sort_order)
        )
        return [(v.key, v.label) for v in rows]


CHARGE, CHARGE_LABEL = choices("products__charge_type")[0]
BASIS, BASIS_LABEL = choices("products__pricing_basis")[0]
UNIT, UNIT_LABEL = choices("products__unit_of_measure")[0]
SEGMENT_LABEL = dict(choices("theme"))


def product_body(suffix: str, **extra) -> dict:
    """Everything the register asks of a product, as the form would send it."""
    return {
        "sku_code": f"{SKU}-{suffix}",
        "item_name": f"Products probe {suffix}",
        "catalogue_category": "SEGMENT_LIBRARY",
        "description": "Made by test_products_editing.py.",
        "charge_type": CHARGE,
        "pricing_basis": BASIS,
        "unit_of_measure": UNIT,
        "supports_perpetual": False,
        "quantity_allowed": False,
        **extra,
    }


def new_product(suffix: str, **extra):
    r = client.post("/api/products", json=product_body(suffix, **extra))
    if r.status_code == 201:
        created["products"].append(r.json()["id"])
    return r


def xlsx(sheet: str, rows: list[list]) -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def upload(body: bytes, *, commit: bool):
    path = "/api/spreadsheets/products/import" + ("" if commit else "/preview")
    return client.post(path, params={"filename": "products.xlsx"}, content=body, headers={"Content-Type": "application/octet-stream"})


def number_of(product_id: str) -> int:
    return int(product_id.split("-")[1])


with SessionLocal() as db:
    if db.get(User, VIEWER) is None:
        stamp = now_utc()
        db.add(User(user_id=VIEWER, name="Products Viewer", email="usr-t-pr1@example.test", active=True, created_at=stamp, updated_at=stamp))
        db.flush()
        db.add(UserRole(user_id=VIEWER, role_id="VIEWER"))
        db.commit()

try:
    print("=" * 74)
    print("  PRODUCTS — the markets a product serves, and changing the catalogue")
    print("=" * 74)

    catalogue = client.get("/api/products").json()
    if not catalogue:
        raise SystemExit("The catalogue is not loaded — run load_price_book.py --apply, then run_tests.py again.")

    print("\n1  A product is made, and asked only for a product's fields")
    r = client.post("/api/products", json={"description": "no name, no SKU"})
    check(
        "a product with no name, SKU or category is refused, each named",
        r.status_code == 422 and code_of(r) == "REQUIRED_FIELDS_MISSING" and missing_of(r) == {"sku_code", "item_name", "catalogue_category"},
        f"{r.status_code} {r.text[:300]}",
    )
    r = client.post("/api/products", json={"sku_code": f"{SKU}-BARE", "item_name": "Products probe bare", "catalogue_category": "SEGMENT_LIBRARY"})
    bare_missing = missing_of(r)
    if r.status_code == 201:
        created["products"].append(r.json()["id"])
    check("the register's other required fields are asked for too", r.status_code == 422 and "unit_of_measure" in bare_missing, f"{r.status_code} {r.text[:300]}")
    check(
        "…but never the price matrix's or a pricing table's",
        not bare_missing & {"product", "size", "list_price_usa_usd", "price_book_version", "size_code", "region_code"},
        str(sorted(bare_missing)),
    )
    check("…and a box left unticked is an answer, not a gap", not bare_missing & {"supports_perpetual", "quantity_allowed"}, str(sorted(bare_missing)))

    r = new_product("ONE")
    check("a complete product is created", r.status_code == 201, f"{r.status_code} {r.text[:300]}")
    if r.status_code != 201:
        raise SystemExit("A product could not be created — nothing after this can be proved.")
    one = r.json()
    ONE = one["id"]
    check("it is given the next product number", ONE.startswith("PRD-") and ONE not in {p["id"] for p in catalogue}, ONE)
    check("it starts active, with no list price yet", one["active"] is True and one["prices"] == [], str(one)[:200])
    check("the boxes left unticked read No", one["has_annual_updates_option"] is False and one["supports_perpetual"] is False)
    check("it is on the list", ONE in {p["id"] for p in client.get("/api/products").json()})
    with SessionLocal() as db:
        audit = db.scalars(select(AuditLog).where(AuditLog.record_module == "products", AuditLog.record_id == ONE)).all()
    check("who created it is recorded", [a.action for a in audit] == ["created"] and audit[0].actor == USER.user_id, str([(a.action, a.actor) for a in audit]))

    r = new_product("one")
    check("the same SKU Code, in any case, is refused", r.status_code == 409 and code_of(r) == "SKU_EXISTS", f"{r.status_code} {r.text[:200]}")
    r = new_product("TWO", effective_from="2026-10-01", effective_to="2026-09-01")
    check("an end before its start is refused", r.status_code == 422 and code_of(r) == "EFFECTIVE_DATES", f"{r.status_code} {r.text[:200]}")
    r = new_product("TWO", effective_from="2026-10-01", effective_to="2027-09-30")
    TWO = r.json()["id"] if r.status_code == 201 else ""
    check("a second product is created", r.status_code == 201 and number_of(TWO) > number_of(ONE), r.text[:200])

    print("\n2  A product is edited")
    r = client.put(f"/api/products/{ONE}", json={**one, "item_name": "Products probe one, renamed", "display_order": 7})
    check("the form's whole record saves", r.status_code == 200 and r.json()["item_name"] == "Products probe one, renamed" and r.json()["display_order"] == 7, r.text[:300])
    r = client.patch(f"/api/products/{ONE}", json={"sku_code": f"{SKU}-two"})
    check("it can't take another product's SKU Code", r.status_code == 409 and code_of(r) == "SKU_EXISTS", f"{r.status_code} {r.text[:200]}")
    r = client.patch(f"/api/products/{ONE}", json={"sku_code": f"  {SKU}-ONE  "})
    check("its own SKU Code, re-sent with spaces, is fine and trimmed", r.status_code == 200 and r.json()["sku_code"] == f"{SKU}-ONE", r.text[:200])
    r = client.patch(f"/api/products/{ONE}", json={"item_name": "  "})
    check("its name can't be emptied", r.status_code == 422 and missing_of(r) == {"item_name"}, f"{r.status_code} {r.text[:200]}")
    r = client.patch(f"/api/products/{ONE}", json={"unit_of_measure": None})
    check("nor can another required field", r.status_code == 422 and missing_of(r) == {"unit_of_measure"}, f"{r.status_code} {r.text[:200]}")
    r = client.patch(f"/api/products/{TWO}", json={"effective_to": "2026-01-01"})
    check("an edit can't end it before it starts", r.status_code == 422 and code_of(r) == "EFFECTIVE_DATES", f"{r.status_code} {r.text[:200]}")
    r = client.patch(f"/api/products/{ONE}", json={"list_price_usa_usd": 5, "size": "XS", "rfp_documnet": "x"})
    check("a key that is no field of a product is not stored", r.status_code == 200 and "rfp_documnet" not in r.json() and "size" not in r.json(), r.text[:200])
    check("a product not there is a 404", client.patch("/api/products/PRD-9999", json={"item_name": "x"}).status_code == 404)
    with SessionLocal() as db:
        last = db.scalars(
            select(AuditLog).where(AuditLog.record_module == "products", AuditLog.record_id == ONE).order_by(AuditLog.id.desc())
        ).first()
    check("an edit is recorded with what moved", last is not None and last.action == "updated", str(last.action if last else None))

    print("\n3  Segment -> Segment Category -> Subsegment")
    deps = client.get("/api/spec").json().get("dependencies", {}).get("products", {})
    mapped = deps.get("segment_categories", {}).get("parent") == "segments" and deps.get("subsegments", {}).get("parent") == "segment_categories"
    check("the register publishes the two maps", mapped, str(list(deps)))
    if mapped:
        cat_map, sub_map = deps["segment_categories"]["map"], deps["subsegments"]["map"]
        seg_a, seg_b = sorted(cat_map)[:2]
        cat_a, cat_b = cat_map[seg_a][0], cat_map[seg_b][0]
        sub_a, sub_b = sub_map[cat_a][0], sub_map[cat_b][0]
        check("41 categories and 132 subsegments are mapped", (sum(map(len, cat_map.values())), sum(map(len, sub_map.values()))) == (41, 132))
        check("every one of the seven Segments has categories", len(cat_map) == 7, str(sorted(cat_map)))
        r = client.patch(f"/api/products/{ONE}", json={"segments": [seg_a, seg_b], "segment_categories": [cat_a, cat_b], "subsegments": [sub_a, sub_b]})
        got = r.json() if r.status_code == 200 else {}
        check(
            "a product takes several of each, across two Segments",
            (got.get("segments"), got.get("segment_categories"), got.get("subsegments")) == ([seg_a, seg_b], [cat_a, cat_b], [sub_a, sub_b]),
            f"{r.status_code} {r.text[:300]}",
        )
        check("…and reads them back", client.get(f"/api/products/{ONE}").json().get("subsegments") == [sub_a, sub_b])
        rows = client.get("/api/products", params={"segments": seg_b}).json()
        check("the list filters by Segment", ONE in {p["id"] for p in rows} and TWO not in {p["id"] for p in rows}, str([p["id"] for p in rows]))
        r = client.patch(f"/api/products/{TWO}", json={"segments": [seg_a], "segment_categories": [cat_a, cat_b]})
        check("a category outside the Segments chosen is refused", r.status_code == 422 and code_of(r) == "DEPENDENT_CHOICE", f"{r.status_code} {r.text[:200]}")
        r = client.patch(f"/api/products/{TWO}", json={"segment_categories": [cat_a]})
        check("a category with no Segment is refused", r.status_code == 422 and code_of(r) == "DEPENDENT_CHOICE", f"{r.status_code} {r.text[:200]}")
        r = client.patch(f"/api/products/{TWO}", json={"segments": [seg_a], "segment_categories": [cat_a], "subsegments": [sub_b]})
        check("a subsegment outside its categories is refused", r.status_code == 422 and code_of(r) == "DEPENDENT_CHOICE", f"{r.status_code} {r.text[:200]}")
        r = client.patch(f"/api/products/{ONE}", json={"item_name": "Products probe one"})
        check("an edit elsewhere leaves them alone", r.status_code == 200 and r.json().get("segments") == [seg_a, seg_b], r.text[:200])
        r = client.patch(f"/api/products/{ONE}", json={"segments": [], "segment_categories": [], "subsegments": []})
        check("they are optional, and can be cleared", r.status_code == 200 and r.json().get("segments") == [], r.text[:200])

    print("\n4  Deactivate, and delete")
    requirements.ENFORCED = False  # the Lead below is made with a handful of fields
    account = client.post("/api/accounts", json={"account_name": "Products editing probe — client", "account_type": ["END_CLIENT"]}).json()["id"]
    created["accounts"].append(account)
    r = client.post(
        "/api/leads",
        json={"opportunity_name": "Products editing probe", "project_stage": "0_CONNECT", "lead_status": "OPEN", "currency": "USD", "end_client": account, "sap_solution_suite": ONE, "secondary_sap_suites": [TWO]},
    )
    lead = r.json().get("id") if r.status_code == 201 else None
    if lead:
        created["leads"].append(lead)
    check("a Lead picks the new product", r.status_code == 201, r.text[:300])
    requirements.ENFORCED = True

    r = client.patch(f"/api/products/{ONE}/active", json={"active": False})
    check("a product is deactivated", r.status_code == 200 and r.json()["active"] is False, r.text[:200])
    check("…and still reads on the Lead naming it", (client.get(f"/api/leads/{lead}").json().get("__labels") or {}).get("sap_solution_suite") == "Products probe one")
    r = client.patch(f"/api/leads/{lead}", json={"suite_demonstrated": ONE})
    check("…but can't be newly picked", r.status_code == 422 and code_of(r) == "PRODUCT_INACTIVE", f"{r.status_code} {r.text[:200]}")
    r = client.patch(f"/api/products/{ONE}/active", json={"active": True})
    check("it is reactivated", r.status_code == 200 and r.json()["active"] is True)

    r = client.delete(f"/api/products/{ONE}")
    detail = r.json().get("detail", {}) if r.status_code == 409 else {}
    check("a product a Lead names can't be deleted", r.status_code == 409 and code_of(r) == "PRODUCT_IN_USE" and detail.get("records") == [lead], f"{r.status_code} {r.text[:200]}")
    r = client.delete(f"/api/products/{TWO}")
    check("…nor one named in a several-pick", r.status_code == 409 and code_of(r) == "PRODUCT_IN_USE", f"{r.status_code} {r.text[:200]}")
    check("…and it is still there", client.get(f"/api/products/{ONE}").status_code == 200)
    r = client.patch(f"/api/leads/{lead}", json={"sap_solution_suite": None, "secondary_sap_suites": []})
    check("taken off the Lead", r.status_code == 200, r.text[:200])
    with SessionLocal() as db:
        book = db.scalar(select(PriceBook.price_book_id).where(PriceBook.is_standard.is_(True)))
        edition = db.scalar(select(PriceBookEdition.edition_id).where(PriceBookEdition.status == "CURRENT"))
        db.add(PriceBookEntry(price_book_id=book, edition_id=edition, product_id=TWO, size_id=None, list_price=100))
        db.commit()
    r = client.delete(f"/api/products/{TWO}")
    check("it is deleted for good", r.status_code == 204, f"{r.status_code} {r.text[:200]}")
    check("…a 404 from then on", client.get(f"/api/products/{TWO}").status_code == 404 and client.delete(f"/api/products/{TWO}").status_code == 404)
    with SessionLocal() as db:
        check("…its list prices with it", db.scalar(select(func.count()).select_from(PriceBookEntry).where(PriceBookEntry.product_id == TWO)) == 0)
        gone = db.scalars(select(AuditLog).where(AuditLog.record_module == "products", AuditLog.record_id == TWO, AuditLog.action == "deleted")).all()
    check("…and who deleted it is recorded", len(gone) == 1 and gone[0].actor == USER.user_id)
    r = new_product("THREE")
    check("its number is never used again", r.status_code == 201 and number_of(r.json()["id"]) > number_of(TWO), r.text[:120])

    print("\n5  Only an admin or a developer changes the catalogue")
    probe = xlsx("Products", [["Item Name", "SKU Code"], ["x", f"{SKU}-V"]])
    act_as_viewer()
    check("anyone with a role reads it", client.get("/api/products").status_code == 200 and client.get(f"/api/products/{ONE}").status_code == 200)
    for label, r in (
        ("create", client.post("/api/products", json=product_body("VIEWER"))),
        ("edit", client.put(f"/api/products/{ONE}", json={"item_name": "x"})),
        ("patch", client.patch(f"/api/products/{ONE}", json={"item_name": "x"})),
        ("deactivate", client.patch(f"/api/products/{ONE}/active", json={"active": False})),
        ("delete", client.delete(f"/api/products/{ONE}")),
        ("import", upload(probe, commit=True)),
        ("preview an import", upload(probe, commit=False)),
    ):
        check(f"a Viewer cannot {label}", r.status_code == 403 and code_of(r) == "CATALOGUE_EDITOR_ONLY", f"{r.status_code} {r.text[:200]}")
    check("…and nothing changed", client.get(f"/api/products/{ONE}").json()["item_name"] == "Products probe one")
    check("a Viewer can export", client.get("/api/spreadsheets/products/export").status_code == 200)
    USER = sign_in_as_admin(app)

    print("\n6  Import and export")
    r = client.get("/api/spreadsheets/products/template")
    labels = [str(c.value).replace(" *", "") for c in load_workbook(io.BytesIO(r.content)).worksheets[1][3] if c.value] if r.status_code == 200 else []
    check("the sample file has a product's columns", r.status_code == 200 and {"Item Name", "SKU Code", "Catalogue Category", "Segment"} <= set(labels), f"{r.status_code} {labels}")
    check("…and none of the price matrix's or the pricing tables'", not {"Size", "List Price (USA USD)", "Price Book Version", "Size Code", "Region Code"} & set(labels), str(labels))
    header = ["Item Name", "SKU Code", "Catalogue Category", "Description", "Charge Type", "Pricing Basis", "Unit of Measure", "Segment"]
    two_segments = f"{SEGMENT_LABEL['SMART_CITIES']}, {SEGMENT_LABEL['DATA_CENTERS']}"
    good = ["Products probe sheet", f"{SKU}-SHEET", "Segment Library", "From a sheet.", CHARGE_LABEL, BASIS_LABEL, UNIT_LABEL, two_segments]
    dupe = ["Products probe dupe", f"{SKU}-ONE", "Segment Library", "Same SKU.", CHARGE_LABEL, BASIS_LABEL, UNIT_LABEL, None]
    r = upload(xlsx("Products", [header, good, dupe]), commit=False)
    report = r.json() if r.status_code == 200 else {}
    check("a preview reads the rows and saves nothing", r.status_code == 200 and report.get("total") == 2 and report.get("ready") == 1, f"{r.status_code} {str(report)[:400]}")
    check("the row with a SKU Code already used is the one refused", [e["row"] for e in report.get("errors", [])] == [3], str(report.get("errors"))[:300])
    check("…nothing saved by the preview", f"{SKU}-SHEET" not in {p["sku_code"] for p in client.get("/api/products").json()})
    r = upload(xlsx("Products", [header, good]), commit=True)
    made = r.json().get("created", []) if r.status_code == 200 else []
    created["products"].extend(made)
    check("the good row imports", r.status_code == 200 and len(made) == 1, f"{r.status_code} {r.text[:300]}")
    if made:
        sheet_row = client.get(f"/api/products/{made[0]}").json()
        check("with its Segments read by name", sheet_row.get("segments") == ["SMART_CITIES", "DATA_CENTERS"], str(sheet_row.get("segments")))
        check("…and its unticked boxes as No", sheet_row["supports_perpetual"] is False and sheet_row["active"] is True)
    r = client.get("/api/spreadsheets/products/export", params={"sku_code_starts": SKU})
    names = load_workbook(io.BytesIO(r.content)).sheetnames if r.status_code == 200 else []
    check(
        "the export is a workbook of the product's own sheets",
        r.status_code == 200 and "Summary" in names and not any(word in n.lower() for n in names for word in ("price", "size", "region", "rate", "support")),
        str(names),
    )
    if r.status_code == 200:
        summary = [[str(c) for c in row] for row in load_workbook(io.BytesIO(r.content))["Summary"].iter_rows(values_only=True)]
        check(
            "…holding what the list shows under the same filter",
            any(f"{SKU}-ONE" in row for row in summary) and not any(catalogue[0]["sku_code"] in row for row in summary),
            str(summary[:4])[:300],
        )

finally:
    from app.auth import current_user as _current_user_dep

    sign_in_as_admin(app)
    requirements.ENFORCED = False
    for lead_id in created["leads"]:
        client.delete(f"/api/leads/{lead_id}")
    for account_id in created["accounts"]:
        client.delete(f"/api/accounts/{account_id}")
    with SessionLocal() as db:
        ids = list(db.scalars(select(Product.product_id).where(Product.sku_code.like(f"{SKU}%"))))
        db.execute(delete(Product).where(Product.product_id.in_(ids)))
        db.execute(delete(UserRole).where(UserRole.user_id == VIEWER))
        db.execute(delete(User).where(User.user_id == VIEWER))
        db.commit()
        leftover = list(db.scalars(select(Product.sku_code).where(Product.sku_code.like(f"{SKU}%"))))
    app.dependency_overrides.pop(_current_user_dep, None)
    print("\nCleanup", "done" if not leftover else f"left {leftover}")

print()
if failures:
    print(f"FAILED — {len(failures)}:")
    for failure in failures:
        print(f"  · {failure}")
    sys.exit(1)
print("All product editing checks passed.")
