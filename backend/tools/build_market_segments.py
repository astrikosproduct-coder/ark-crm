"""
Write backend/data/market_segments.json — Segment -> Segment Category -> Subsegment.

    python tools/build_market_segments.py
    python tools/build_market_segments.py "C:/path/to/another copy.xlsx"

Reads the go-to-market taxonomy workbook, docs/Segments_and_Subsegments.xlsx
(sheet "Subsegments": #, Segment, Category, Subsegment, Examples / Notes,
Priority, Owner). The file it writes is COMMITTED and is what
products_market_segments_metadata.py loads into the register, so production
needs neither the workbook nor this script. Re-run it only to take a newer
copy of the workbook.

WHAT IS TAKEN, AND WHAT IS NOT
------------------------------
Segment     The seven are ARK's existing Segment list (picklist `theme`, used
            by Leads and Accounts since Release 1.1). They are matched to its
            choices by the table below — written out, never guessed — so a
            product's Segment and a Lead's Segment are the same choice. Three
            are spelled differently in the workbook; ARK's labels are KEPT, and
            the difference is printed by the metadata script. Renaming one is
            a change to Leads and Accounts too, made in Setup.
Category    41, each under one Segment.
Subsegment  132, each under one Category. Names are unique across the sheet,
            which is checked here: a repeat would need a longer key.
Examples    Kept in the JSON beside each Subsegment. A dropdown choice has
            nowhere to show it yet.
Priority, Owner   Both columns are empty in the workbook, and a dropdown choice
            has nowhere to hold them. Not taken.

Keys are the name in capitals with underscores (PUBLIC_SAFETY), the form every
other ARK choice has. They are what a record stores, so they never change once
loaded; the label beside one can be renamed in Setup.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT.parent / "docs" / "Segments_and_Subsegments.xlsx"
OUT = ROOT / "data" / "market_segments.json"
SHEET = "Subsegments"
HEADERS = ("Segment", "Category", "Subsegment")

#: The workbook's Segment -> the key of ARK's Segment choice (picklist `theme`).
SEGMENT_KEYS = {
    "Smart Cities": "SMART_CITIES",
    "Smart Buildings & Campuses": "SMART_BUILDINGS_CAMPUSES",
    "Smart Transport & Infrastructure": "SMART_TRANSPORT_INFRASTRUCTURE",
    "Industry 4.0 & Manufacturing": "INDUSTRY_4_0_MANUFACTURING",
    "Energy, Utilities & Critical Infra": "ENERGY_UTILITIES_CRITICAL_INFRASTRUCTURE",
    "Data Centres": "DATA_CENTERS",
    "Defence, Space & Protected Markets": "DEFENCE_PROTECTED_MARKETS",
}


def key_of(label: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", label.upper()).strip("_")


def main() -> int:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE
    if not source.exists():
        raise SystemExit(f"STOPPED: {source} is not there.")
    book = load_workbook(source, data_only=True, read_only=True)
    if SHEET not in book.sheetnames:
        raise SystemExit(f"STOPPED: {source.name} has no sheet named {SHEET!r}.")
    rows = [list(r) for r in book[SHEET].iter_rows(values_only=True)]
    header = [str(c or "").strip() for c in rows[0]]
    missing = [h for h in HEADERS if h not in header]
    if missing:
        raise SystemExit(f"STOPPED: the {SHEET} sheet has no column {missing}.")
    at = {name: header.index(name) for name in HEADERS}
    notes_at = next((i for i, h in enumerate(header) if h.lower().startswith("examples")), None)

    segments: list[dict] = []
    categories: list[dict] = []
    subsegments: list[dict] = []
    seen_category: dict[str, str] = {}
    seen_sub: set[str] = set()
    for number, row in enumerate(rows[1:], start=2):
        segment, category, sub = (str(row[at[h]] or "").strip() for h in HEADERS)
        if not (segment or category or sub):
            continue
        if not (segment and category and sub):
            raise SystemExit(f"STOPPED: row {number} is missing its Segment, Category or Subsegment.")
        if segment not in SEGMENT_KEYS:
            raise SystemExit(
                f"STOPPED: row {number} names the Segment {segment!r}, which is not one of ARK's seven. "
                "Add it to SEGMENT_KEYS with the key of the choice it is."
            )
        segment_key = SEGMENT_KEYS[segment]
        if not any(s["key"] == segment_key for s in segments):
            segments.append({"key": segment_key, "workbook_label": segment})

        category_key = key_of(category)
        if category_key not in seen_category:
            seen_category[category_key] = segment_key
            categories.append({"key": category_key, "label": category, "segment": segment_key})
        elif seen_category[category_key] != segment_key:
            raise SystemExit(f"STOPPED: the Category {category!r} (row {number}) sits under two Segments.")

        sub_key = key_of(sub)
        if sub_key in seen_sub:
            raise SystemExit(f"STOPPED: the Subsegment {sub!r} (row {number}) appears twice.")
        seen_sub.add(sub_key)
        examples = str(row[notes_at] or "").strip() if notes_at is not None and notes_at < len(row) else ""
        subsegments.append({"key": sub_key, "label": sub, "category": category_key, "examples": examples})

    document = {
        "source": source.name,
        "segments": segments,
        "categories": categories,
        "subsegments": subsegments,
    }
    OUT.write_text(json.dumps(document, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {OUT}")
    print(f"  {len(segments)} segments, {len(categories)} categories, {len(subsegments)} subsegments")
    return 0


if __name__ == "__main__":
    sys.exit(main())
