"""
Excel import and export — the sample file, the preview, all-or-nothing commit,
and an export that matches the list screen.

    python run_tests.py test_spreadsheets.py

UX roadmap item 4, 17 Sep 2026 — app/spreadsheets/, app/routers/spreadsheets.py.
"""

import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy import func, select  # noqa: E402

import test_db  # noqa: E402

test_db.require_test_database()

from app.database import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Lead  # noqa: E402
from app.spreadsheets.fields import CellError, Column, read_cell  # noqa: E402
from test_support import sign_in_as_admin  # noqa: E402

client = TestClient(app)
USER = sign_in_as_admin(app)
failures: list[str] = []
created: dict[str, list[str]] = {"leads": [], "contacts": [], "accounts": []}
TAG = "Sheet probe"


def check(label: str, condition: bool, detail: str = "") -> None:
    if condition:
        print(f"  pass  {label}")
    else:
        failures.append(f"{label}{' — ' + detail if detail else ''}")
        print(f"  FAIL  {label}{' — ' + detail if detail else ''}")


def code_of(response) -> str | None:
    try:
        detail = response.json().get("detail")
    except ValueError:
        return None
    return detail.get("code") if isinstance(detail, dict) else None


def xlsx(sheet: str, rows: list[list]) -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def filled_sample(module: str, rows: list[dict]) -> bytes:
    """ARK's own sample file, filled in from row 6 by label, as a person would."""
    r = client.get(f"/api/spreadsheets/{module}/template")
    book = load_workbook(io.BytesIO(r.content))
    ws = book[book.sheetnames[1]]
    headers = {str(c.value).replace(" *", ""): c.column for c in ws[3] if c.value}
    for offset, values in enumerate(rows):
        for label, value in values.items():
            ws.cell(row=6 + offset, column=headers[label], value=value)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def lead_count() -> int:
    with SessionLocal() as db:
        return db.scalar(select(func.count()).select_from(Lead))


def upload(module: str, body: bytes, filename: str, *, commit: bool, mapping: str | None = None, answers: str | None = None):
    path = f"/api/spreadsheets/{module}/import" + ("" if commit else "/preview")
    params = {"filename": filename}
    if mapping:
        params["mapping"] = mapping
    if answers:
        params["answers"] = answers
    return client.post(path, params=params, content=body, headers={"Content-Type": "application/octet-stream"})


try:
    print("=" * 74)
    print("  SPREADSHEETS — sample, preview, commit, export")
    print("=" * 74)

    end_client = client.post("/api/accounts", json={"account_name": f"{TAG} client", "account_type": ["END_CLIENT"]}).json()["id"]
    created["accounts"].append(end_client)
    other_client = client.post("/api/accounts", json={"account_name": f"{TAG} other client", "account_type": ["END_CLIENT"]}).json()["id"]
    created["accounts"].append(other_client)

    print("\n1  The sample file")
    r = client.get("/api/spreadsheets/leads/template")
    check("the Leads sample downloads as XLSX, named for Astrikos", r.status_code == 200 and "spreadsheetml" in r.headers.get("content-type", "") and "Astrikos" in r.headers.get("content-disposition", ""), str(r.headers.get("content-disposition")))
    book = load_workbook(io.BytesIO(r.content))
    check("…Instructions and Leads visible, Lists hidden", book.sheetnames[:2] == ["Instructions", "Leads"] and book["Lists"].sheet_state == "hidden", str(book.sheetnames))
    ws = book["Leads"]
    check("…row 1 is the Astrikos band on every sheet", ws["A1"].value == "ASTRIKOS" and book["Instructions"]["A1"].value == "ASTRIKOS")
    headers = [c.value for c in ws[3]]
    check("…labels on row 3, the name first and marked required", headers[0] == "Opportunity Name *" and "opportunity_name" not in headers, str(headers[:6]))
    check("…section bands on row 2", "Stage 0" in str(ws["A2"].value), str(ws["A2"].value))
    check("…no system or stage-set columns", not {"Created By", "Lead Status", "Probability (%)"} & set(headers), str(headers))
    check("…no Stage 1+ fields", "Demo Date" not in headers and "Interest Level" not in headers)
    hints = dict(zip(headers, [c.value for c in ws[4]]))
    conditions = dict(zip(headers, [c.value for c in ws[5]]))
    check("…row 4 says a lookup must be picked from ARK", hints.get("End Client") == "Pick from ARK · must already exist", str(hints.get("End Client")))
    check("…row 5 says when a column applies", conditions.get("Incremental Value") == "Only if Opportunity Type is Expansion", str(conditions.get("Incremental Value")))
    letter = {h: c.column_letter for h, c in zip(headers, ws[3])}
    by_letter = {str(dv.sqref).split(":")[0].rstrip("0123456789"): dv for dv in ws.data_validations.dataValidation}
    end_client_dv = by_letter.get(letter.get("End Client"))
    check("…End Client is a dropdown of live account names", end_client_dv is not None and end_client_dv.type == "list" and "Lists!" in str(end_client_dv.formula1), str(end_client_dv and end_client_dv.formula1))
    names = []
    if end_client_dv is not None:
        list_letter = str(end_client_dv.formula1).split("!")[1].split(":")[0].replace("$", "").rstrip("0123456789")
        names = [c.value for c in book["Lists"][list_letter]]
    check("…listing End Client accounts", f"{TAG} client" in names, str(names[:8]))
    check("…with the hint shown when a cell is selected", end_client_dv is not None and "Pick from ARK" in str(end_client_dv.prompt))
    r = client.get("/api/spreadsheets/leads/template", params={"format": "csv"})
    check("the CSV sample is the header row", r.status_code == 200 and r.content.decode("utf-8-sig").startswith("Opportunity Name *"), r.content[:120].decode("utf-8-sig", "replace"))
    r = client.get("/api/spreadsheets/opportunities/template")
    check("Opportunities offer no import", r.status_code == 404 and code_of(r) == "IMPORT_NOT_OFFERED", r.text[:200])

    r = client.get("/api/spreadsheets/accounts/template")
    book = load_workbook(io.BytesIO(r.content))
    ws = book["Accounts"]
    headers = [c.value for c in ws[3]]
    conditions = dict(zip(headers, [c.value for c in ws[5]]))
    check("Account Type is one Yes/No column per choice", "Account Type · End Client" in headers and "Account Type · Partner / SI" in headers and "Account Type" not in headers, str(headers[:7]))
    check("…and Partner Tier's condition reads as one sentence", str(conditions.get("Partner Tier", "")).startswith("Only if Account Type includes Partner / SI, "), str(conditions.get("Partner Tier")))

    print("\n2  Preview checks every row and saves nothing")
    before = lead_count()
    bad = xlsx("Leads", [
        ["Opportunity Name *", "End Client", "BD Owner", "Segment", "Estimated Value", "Lead Stage"],
        [f"{TAG} good", f"{TAG} client", USER.name, None, "250,000", None],
        [f"{TAG} bad choice", f"{TAG} other client", None, "Not a segment", None, None],
        [f"{TAG} unknown client", "No such client ever", None, None, "abc", None],
        [None, f"{TAG} client", None, None, None, None],
        [f"{TAG} too far", f"{TAG} other client", None, None, None, "1 Demo"],
        [None, None, None, None, None, None],
    ])
    r = upload("leads", bad, "leads.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    check("the preview answers", r.status_code == 200, r.text[:300])
    check("…counts rows, ignoring blank ones", report.get("total") == 5, str(report.get("total")))
    check("…1 ready, 4 to fix", report.get("ready") == 1 and report.get("failed") == 4, f"{report.get('ready')} / {report.get('failed')}")
    messages = {e["row"]: " | ".join(e["messages"]) for e in report.get("errors", [])}
    check("…a bad choice is named in words", "isn't one of the choices" in messages.get(3, ""), messages.get(3, ""))
    check("…an unknown account and a non-number both reported on the row", "no account called" in messages.get(4, "") and "enter a number" in messages.get(4, ""), messages.get(4, ""))
    check("…a row with no name is refused", "Opportunity Name is required" in messages.get(5, ""), messages.get(5, ""))
    check("…a Stage 1 row is refused with the reason", "Stage 0" in messages.get(6, ""), messages.get(6, ""))
    check("…and nothing was saved", lead_count() == before, f"{before} -> {lead_count()}")

    print("\n3  Commit is all-or-nothing")
    r = upload("leads", bad, "leads.xlsx", commit=True)
    check("committing a file with problems is refused", r.status_code == 422 and code_of(r) == "IMPORT_NOT_READY", r.text[:200])
    check("…and the good row was not saved either", lead_count() == before)

    dupes = xlsx("Leads", [
        ["Opportunity Name *", "End Client", "Estimated Value"],
        [f"{TAG} dup A", f"{TAG} client", 1000],
        [f"{TAG} dup B", f"{TAG} client", 2000],
    ])
    r = upload("leads", dupes, "dupes.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    messages = {e["row"]: " | ".join(e["messages"]) for e in report.get("errors", [])}
    check("two rows for one End Client trip the duplicate guard, as two saves would", "already has an open pursuit" in messages.get(3, ""), str(report.get("errors")))
    check("…and the fix names the reason column", "Not a Duplicate" in messages.get(3, ""), messages.get(3, ""))
    reasoned = xlsx("Leads", [
        ["Opportunity Name *", "End Client", "Not a Duplicate — Reason"],
        [f"{TAG} dup A", f"{TAG} client", None],
        [f"{TAG} dup B", f"{TAG} client", "Separate datacentre tender."],
    ])
    r = upload("leads", reasoned, "reasoned.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    check("…and giving that reason lets the second row through", report.get("ready") == 2 and report.get("failed") == 0, str(report.get("errors")))

    good = xlsx("Leads", [
        ["Opportunity Name *", "Client", "BD Owner", "Estimated Value"],
        [f"{TAG} import one", f"{TAG} client", USER.email, 125000],
        [f"{TAG} import two", f"{TAG} other client", USER.name, None],
    ])
    r = upload("leads", good, "good.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    check("an unrecognised header is reported for mapping", report.get("unmapped") == ["Client"], str(report.get("unmapped")))
    r = upload("leads", good, "good.xlsx", commit=True, mapping='{"Client": "end_client"}')
    report = r.json() if r.status_code == 200 else {}
    created["leads"].extend(report.get("created") or [])
    check("with the column mapped, both rows import", r.status_code == 200 and report.get("committed") and len(report.get("created") or []) == 2, r.text[:300])
    if report.get("created"):
        lead = client.get(f"/api/leads/{report['created'][0]}").json()
        check("…through the create route: stamped by the importer", lead.get("created_by") == USER.user_id, str(lead.get("created_by")))
        check("…at Stage 0, Open, with the End Client resolved by name", lead.get("project_stage", "").startswith("0") and lead.get("lead_status") == "OPEN" and lead.get("end_client") == end_client, str({k: lead.get(k) for k in ("project_stage", "lead_status", "end_client")}))
        check("…and the person matched by email", lead.get("bd_owner") == USER.user_id)
        history = client.get("/api/audit-log", params={"record_id": report["created"][0]}).json()
        check("…with its own History entry", any(h.get("action") == "created" for h in history), str(history)[:200])

    print("\n4  Filling in the sample file")
    body = filled_sample("accounts", [
        {"Account Name": f"{TAG} partner", "Account Type · Partner / SI": "Yes", "Partner Tier": "Registered"},
        {"Account Name": f"{TAG} end client", "Account Type · End Client": "Yes", "Partner Tier": "Registered"},
    ])
    r = upload("accounts", body, "Astrikos ARK - Accounts import sample.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    check("the filled sample reads from row 6, skipping the guidance rows", r.status_code == 200 and report.get("total") == 2 and report.get("header_row") == 3, r.text[:300])
    check("…both rows ready", report.get("ready") == 2 and report.get("failed") == 0, str(report.get("errors")))
    warning = " | ".join(m for w in report.get("warnings", []) for m in w["messages"])
    check("…Partner Tier on an End Client is left out with a warning, not refused", report.get("warned") == 1 and "Partner Tier: left out" in warning and report["warnings"][0]["row"] == 7, str(report.get("warnings")))
    r = upload("accounts", body, "sample.xlsx", commit=True)
    report = r.json() if r.status_code == 200 else {}
    created["accounts"].extend(report.get("created") or [])
    check("…and imports", r.status_code == 200 and len(report.get("created") or []) == 2, r.text[:300])
    if len(report.get("created") or []) == 2:
        partner_row = client.get(f"/api/accounts/{report['created'][0]}").json()
        client_row = client.get(f"/api/accounts/{report['created'][1]}").json()
        check("…Yes columns become the Account Type", partner_row.get("account_type") == ["PARTNER_SI"] and client_row.get("account_type") == ["END_CLIENT"], f"{partner_row.get('account_type')} / {client_row.get('account_type')}")
        check("…the partner keeps its tier; the End Client has none", bool(partner_row.get("partner_tier")) and not client_row.get("partner_tier"), f"{partner_row.get('partner_tier')} / {client_row.get('partner_tier')}")

    typed = xlsx("Accounts", [
        ["Account Name *", "Account Type"],
        [f"{TAG} typed one", "SI"],
        [f"{TAG} typed two", "End client, partner / si"],
        [f"{TAG} typed three", "SI"],
        [f"{TAG} typed four", None],
    ])
    r = upload("accounts", typed, "typed.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    unknown = report.get("unknown_values") or []
    check("an unrecognised choice is listed once, with its row count", len(unknown) == 1 and unknown[0]["value"] == "SI" and unknown[0]["rows"] == 2 and unknown[0]["field"] == "account_type", str(unknown))
    check("…with the choices to pick from", any(c["key"] == "PARTNER_SI" for c in (unknown[0]["choices"] if unknown else [])))
    check("…while commas between real choices are read", report.get("ready") == 2 and report.get("failed") == 2, f"{report.get('ready')} / {report.get('failed')}")
    r = upload("accounts", typed, "typed.xlsx", commit=False, answers='{"account_type": {"SI": "PARTNER_SI"}}')
    report = r.json() if r.status_code == 200 else {}
    check("…and answering it once lets every row through", report.get("ready") == 4 and not report.get("unknown_values"), str(report.get("errors")))
    variants = xlsx("Accounts", [["Account Name *", "Account Type: End Client"], [f"{TAG} variant", "yes"]])
    r = upload("accounts", variants, "variants.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    check("a choice column written 'Account Type: End Client' maps too", report.get("ready") == 1 and not report.get("unmapped"), str(report.get("unmapped")))

    missing = xlsx("Leads", [
        ["Opportunity Name *", "End Client"],
        [f"{TAG} m1", "Missing client A"],
        [f"{TAG} m2", "Missing client A"],
        [f"{TAG} m3", "Missing client B"],
    ])
    r = upload("leads", missing, "missing.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    group = (report.get("missing_records") or [{}])[0]
    check("names missing from ARK are grouped per column", group.get("field") == "end_client" and group.get("count") == 2 and group.get("rows") == 3 and group.get("noun") == "account", str(report.get("missing_records")))

    day = Column(api_name="kickoff_date", label="Kick-off", type="date", section="")
    month = Column(api_name="expected_close_month", label="Close month", type="date", section="")

    def read(column, raw):
        try:
            return read_cell(column, raw, None)
        except CellError as exc:
            return f"error: {exc}"

    check("17-Sep-2026 reads", read(day, "17-Sep-2026") == "2026-09-17")
    check("2026-09-17 reads", read(day, "2026-09-17") == "2026-09-17")
    check("25/03/2026 can only be one date, so it reads", read(day, "25/03/2026") == "2026-03-25")
    ambiguous = str(read(day, "03/04/2026"))
    check("03/04/2026 is refused, naming both readings", ambiguous.startswith("error:") and "03-Apr-2026" in ambiguous and "04-Mar-2026" in ambiguous, ambiguous)
    check("a month reads as its first day", read(month, "Oct-2026") == "2026-10-01" and read(month, "10/2026") == "2026-10-01")
    check("text that isn't a date says how to write one", "17-Sep-2026" in str(read(day, "next week")))

    print("\n5  CSV, contacts and partners")
    csv_body = ("Full Name *,Account,Job Title\n" f"{TAG} person,{TAG} client,Director\n").encode("utf-8")
    r = upload("contacts", csv_body, "contacts.csv", commit=True)
    report = r.json() if r.status_code == 200 else {}
    created["contacts"].extend(report.get("created") or [])
    check("a CSV of contacts imports, account matched by name", r.status_code == 200 and len(report.get("created") or []) == 1, r.text[:300])
    partner = xlsx("Partners", [["Account Name *", "Account Type"], [f"{TAG} not a partner", "End Client"]])
    r = upload("partners", partner, "partners.xlsx", commit=False)
    report = r.json() if r.status_code == 200 else {}
    check("a Partners import refuses an account with no partner type", report.get("failed") == 1 and "partner" in " ".join(report["errors"][0]["messages"]).lower(), str(report.get("errors")))
    r = upload("leads", b"not a spreadsheet", "notes.txt", commit=False)
    check("an unsupported file is refused in words", r.status_code == 422 and code_of(r) == "FILE_REFUSED", r.text[:200])

    print("\n6  Export matches the list")
    r = client.get("/api/spreadsheets/leads/export", params={"q": TAG, "_search": "opportunity_name"})
    check("the Leads export downloads", r.status_code == 200 and "spreadsheetml" in r.headers.get("content-type", ""), str(r.status_code))
    book = load_workbook(io.BytesIO(r.content))
    check("…opening on a Summary sheet, then the form's sections", book.sheetnames[0] == "Summary" and len(book.sheetnames) > 3, str(book.sheetnames))
    summary = [row for row in book["Summary"].iter_rows(values_only=True)]
    check("…the Astrikos band on row 1, labels on row 2", summary[0][0] == "ASTRIKOS", str(summary[0][:2]))
    summary = summary[1:]
    listed = client.get("/api/leads", params={"q": TAG, "_search": "opportunity_name"}).json()
    check("…one row per listed lead, under the same filter", len(summary) - 1 == len(listed) == 2, f"{len(summary) - 1} vs {len(listed)}")
    check("…Record ID and Name first, labels as headers", summary[0][:2] == ("Record ID", "Opportunity Name") and "End Client" in summary[0], str(summary[0]))
    names = {row[1]: row for row in summary[1:]}
    one = names.get(f"{TAG} import one")
    col = summary[0].index("End Client") if "End Client" in summary[0] else None
    check("…lookups exported as names, not ids", one is not None and col is not None and one[col] == f"{TAG} client", str(one))
    r = client.get("/api/spreadsheets/opportunities/export")
    check("Opportunities export (export-only module)", r.status_code == 200, str(r.status_code))
    r = client.get("/api/spreadsheets/deals/export")
    check("Deals export", r.status_code == 200, str(r.status_code))
    r = client.get("/api/spreadsheets/partners/export")
    book = load_workbook(io.BytesIO(r.content)) if r.status_code == 200 else None
    check("Partners export holds partners only", book is not None and all(f"{TAG} not a partner" != row[1] for row in book["Summary"].iter_rows(min_row=3, values_only=True)), str(r.status_code))

finally:
    print("\nCleanup")
    for collection in ("contacts", "leads", "accounts"):
        for record_id in reversed(created[collection]):
            client.delete(f"/api/{collection}/{record_id}")

print()
if failures:
    print(f"{len(failures)} failed")
    for f in failures:
        print(f"  - {f}")
    sys.exit(1)
print("all passed")
