"""
Spreadsheet files in and out — XLSX, XLS and CSV.

Reading turns any of the three into `{sheet name: rows of cell values}` and
nothing more; what a cell MEANS is fields.py's business. Writing builds the two
workbooks a person downloads: the import sample file and an export.

Headers are always field LABELS, never api_names — this is a file for BD, not
for a developer. Import maps a header back to its field by label, so a column
renamed in Administration is picked up by the next sample file automatically.

THE SAMPLE FILE'S LAYOUT (agreed 17 Sep 2026)
---------------------------------------------
    row 1   Astrikos brand band
    row 2   the form's sections, one band across each section's columns
    row 3   field labels — amber must be filled, green is needed later
    row 4   what to enter            ("Date · e.g. 17-Sep-2026", "Pick from ARK")
    row 5   when the column applies  ("Only if Account Type includes Partner / SI")
    row 6+  the rows to import

The importer does not depend on those row numbers. It finds the label row by
matching labels, and skips rows 4-5 because their text is exactly the guidance
this module writes (`is_guidance_row`). A plain file with its headers on row 1
imports the same way.
"""

from __future__ import annotations

import csv
import io
import math
import re
from datetime import date, datetime
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .conditions import ConditionError, _tokens
from .fields import Column, Directory, lookup_choices, section_title, split_choices

MAX_UPLOAD_BYTES = 5 * 1024 * 1024

BRAND = "ASTRIKOS"
BRAND_FILL = PatternFill("solid", fgColor="1B2A4A")
BRAND_FONT = Font(bold=True, size=13, color="FFFFFF")
BRAND_SUB_FONT = Font(size=10, color="D5DEEB")
SECTION_FILL = PatternFill("solid", fgColor="D5DEEB")
SECTION_FONT = Font(bold=True, size=10, color="1B2A4A")
HEADER_FILL = PatternFill("solid", fgColor="EEF1F5")
REQUIRED_FILL = PatternFill("solid", fgColor="FDE9C8")
LATER_FILL = PatternFill("solid", fgColor="E3F2E8")
HEADER_FONT = Font(bold=True)
HINT_FONT = Font(size=9, color="4B5563")
CONDITION_FONT = Font(size=9, italic=True, color="8A5A00")

DATE_FORMAT = "dd-mmm-yyyy"
MONTH_FORMAT = "mmm-yyyy"
MONEY_FORMAT = "#,##0"

LABEL_ROW = 3
HINT_ROW = 4
CONDITION_ROW = 5
FIRST_DATA_ROW = 6
TEMPLATE_ROWS = 500  # how far down the dropdowns and formats reach

#: Excel date serials for 1 Jan 1990 and 31 Dec 2100 — the date check's bounds.
_EARLIEST_SERIAL, _LATEST_SERIAL = 32874, 73415
_NUMBER_BOUND = 10**15


class FileRefusal(ValueError):
    """A file that cannot be read at all. The message is shown as-is."""


# ------------------------------------------------------------------ reading


def read_upload(filename: str, data: bytes) -> dict[str, list[list[Any]]]:
    if not data:
        raise FileRefusal("The file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise FileRefusal("The file is larger than 5 MB. Split it into smaller files.")
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return _read_xlsx(data)
    if name.endswith(".xls"):
        return _read_xls(data)
    if name.endswith(".csv"):
        return {"CSV": _read_csv(data)}
    raise FileRefusal("Use an XLSX, XLS or CSV file.")


def _read_xlsx(data: bytes) -> dict[str, list[list[Any]]]:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # openpyxl raises a zoo of types for a damaged file
        raise FileRefusal("This file can't be opened as an XLSX workbook.") from None
    sheets = {}
    for sheet in book.worksheets:
        if sheet.sheet_state != "visible":
            continue  # the sample file's hidden Lists sheet feeds dropdowns; it is never data
        sheets[sheet.title] = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    return sheets


def _read_xls(data: bytes) -> dict[str, list[list[Any]]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception:
        raise FileRefusal("This file can't be opened as an XLS workbook.") from None
    sheets = {}
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        sheets[sheet.name] = rows
    return sheets


def _read_csv(data: bytes) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise FileRefusal("This CSV isn't in a text encoding ARK can read. Save it as CSV UTF-8.")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[cell if cell != "" else None for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def trim_rows(rows: list[list[Any]]) -> list[list[Any]]:
    """Drop fully blank rows at the end — Excel keeps formatted empty rows."""
    out = list(rows)
    while out and all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in out[-1]):
        out.pop()
    return out


def header_key(text: Any) -> str:
    """A header as a lookup key: the label without its required star, case and spacing."""
    return re.sub(r"\s+", " ", str(text or "").replace("*", "")).strip().lower()


def choice_header(column: Column, choice_label: str) -> str:
    """The header of one choice's Yes/No column: "Account Type · Partner / SI"."""
    return f"{column.label} · {choice_label}"


def choice_key(text: Any) -> str:
    """A choice header as a lookup key — "Account Type: End Client" and "Account Type - End Client" match too."""
    return re.sub(r"\s*[·:|–—]\s*|\s+-\s+", " · ", header_key(text))


# ------------------------------------------------------------------ describing a column


def condition_in_words(expr: str | None, labels: dict[str, str], directory: Directory, picklists: dict[str, str | None]) -> str | None:
    """
    `account_type includes 'Partner / SI' || account_type includes 'OEM'`
    → `Account Type includes Partner / SI or OEM`. None if it won't read.
    """
    if not expr:
        return None
    try:
        tokens = _tokens(expr)
    except ConditionError:
        return None
    clauses: list[tuple[str, str, str]] = []  # (label, phrase, value)
    joiners: list[str] = []
    i = 0
    while i < len(tokens):
        kind, text = tokens[i]
        if kind == "id" and i + 2 < len(tokens) and tokens[i + 2][0] == "str":
            op, literal = tokens[i + 1][1], tokens[i + 2][1]
            label = labels.get(text, text.replace("_", " ").title())
            shown = next((option_label for key, option_label, _ in directory.options(picklists.get(text)) if key == literal), literal)
            phrase = {"==": "is", "!=": "is not", "includes": "includes"}.get(op)
            if phrase is None:
                return None
            if literal == "" and op == "!=":
                clauses.append((label, "is filled in", ""))
            else:
                clauses.append((label, phrase, shown))
            i += 3
        elif kind == "op" and text in ("&&", "||"):
            joiners.append("and" if text == "&&" else "or")
            i += 1
        elif kind == "op" and text in ("(", ")"):
            i += 1
        else:
            return None
    if not clauses or len(joiners) != len(clauses) - 1:
        return None
    # "A includes X or A includes Y or A includes Z" reads as "A includes X, Y or Z".
    groups: list[tuple[str, str, list[str]]] = [(clauses[0][0], clauses[0][1], [clauses[0][2]])]
    group_joiners: list[str] = []
    for joiner, (label, phrase, value) in zip(joiners, clauses[1:]):
        last = groups[-1]
        if joiner == "or" and last[0] == label and last[1] == phrase and phrase in ("is", "includes"):
            last[2].append(value)
        else:
            groups.append((label, phrase, [value]))
            group_joiners.append(joiner)
    words = []
    for index, (label, phrase, values) in enumerate(groups):
        if index:
            words.append(group_joiners[index - 1])
        listed = values[0] if len(values) == 1 else ", ".join(values[:-1]) + " or " + values[-1]
        words.append(f"{label} {phrase} {listed}".rstrip())
    return " ".join(words)


def how_to_fill(column: Column, *, choice: bool = False) -> str:
    """Row 4 of the sample file. Short, and the same words for the same kind of field everywhere."""
    if choice:
        return "Yes if it applies"
    kind = column.type
    if kind in ("date", "datetime"):
        return "Month · e.g. Oct-2026" if column.api_name.endswith("_month") else "Date · e.g. 17-Sep-2026"
    if kind == "text" and column.max_length:
        return f"Text · up to {column.max_length} characters"
    if kind == "lookup":
        return "Pick a person in ARK" if column.lookup_target == "user" else "Pick from ARK · must already exist"
    return {
        "text": "Text",
        "longtext": "Text",
        "richtext": "Text",
        "url": "Web address",
        "email": "Email address",
        "phone": "Phone, with country code",
        "number": "Number",
        "currency": "Amount · no currency symbol",
        "percent": "Percent · e.g. 40",
        "checkbox": "Yes or No",
        "picklist": "Dropdown",
        "multiselect": "Choices · separate with ;",
    }.get(kind, "Text")


def condition_note(text: str) -> str:
    """Row 5 of the sample file."""
    return f"Only if {text}"


_CONDITION_PREFIXES = ("only if ", "only when ")


def is_guidance_row(row: list[Any], expected: dict[int, str]) -> bool:
    """
    Row 4 or 5 of a sample file, left in place above the data.

    `expected` is the row-4 text each mapped column would carry. A row is
    guidance when at least one mapped cell is filled and every filled mapped cell
    is either that exact text or a row-5 "Only if …" note. A real record never
    reads like that — its name cell alone would not match.
    """
    filled = 0
    for index, wanted in expected.items():
        cell = row[index] if index < len(row) else None
        if cell is None or (isinstance(cell, str) and not cell.strip()):
            continue
        text = str(cell).strip()
        if text != wanted and not text.lower().startswith(_CONDITION_PREFIXES):
            return False
        filled += 1
    return filled > 0


# ------------------------------------------------------------------ writing


def _brand_row(sheet, subtitle: str, width: int) -> None:
    """Row 1 of every sheet ARK writes: the Astrikos band."""
    sheet.cell(row=1, column=1, value=BRAND).font = BRAND_FONT
    sheet.cell(row=1, column=2 if width > 1 else 1, value=subtitle if width > 1 else f"{BRAND}  ·  {subtitle}")
    if width <= 1:
        sheet.cell(row=1, column=1).font = BRAND_FONT
    else:
        sheet.cell(row=1, column=2).font = BRAND_SUB_FONT
    for index in range(1, max(width, 2) + 1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = BRAND_FILL
        cell.alignment = Alignment(vertical="center")
    if width > 2:
        sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=width)
    sheet.row_dimensions[1].height = 24


def safe_sheet_title(title: str, taken: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", title).strip()[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in taken:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    taken.add(candidate.lower())
    return candidate


def _lines(text: str, width: float) -> int:
    return max(1, math.ceil(len(text) / max(1.0, width * 1.15)))


def _as_text(cell) -> None:
    """openpyxl reads any text starting "=" as a formula. A name or a note must stay text."""
    if cell.data_type == "f":
        cell.data_type = "s"


def template_workbook(
    *,
    plural: str,
    noun: str,
    columns: list[Column],
    required: set[str],
    later: set[str],
    conditions: dict[str, str],
    directory: Directory,
    notes: list[str],
    generated: date,
) -> bytes:
    """
    The sample file: the Instructions sheet, the sheet to fill in (layout in the
    module docstring), and a hidden Lists sheet feeding the dropdowns.
    """
    book = Workbook()
    book.properties.creator = "Astrikos · ARK CRM"
    book.properties.title = f"ARK CRM {plural} import sample"
    taken: set[str] = set()
    instructions = book.active
    instructions.title = safe_sheet_title("Instructions", taken)
    main = book.create_sheet(safe_sheet_title(plural, taken))
    lists = book.create_sheet(safe_sheet_title("Lists", taken))
    list_ranges: dict[str, str] = {}

    def list_range(cache_key: str, title: str, values: list[str]) -> str | None:
        if not values:
            return None
        if cache_key not in list_ranges:
            index = len(list_ranges) + 1
            letter = get_column_letter(index)
            lists.cell(row=1, column=index, value=title).font = HEADER_FONT
            for row, value in enumerate(values, start=2):
                _as_text(lists.cell(row=row, column=index, value=value))
            list_ranges[cache_key] = f"Lists!${letter}$2:${letter}${len(values) + 1}"
        return list_ranges[cache_key]

    # ---- the columns, a multi-choice field spread into one Yes/No column per choice
    slots: list[tuple[Column, str | None]] = []
    for column in columns:
        choices = split_choices(column, directory)
        if choices:
            slots.extend((column, label) for _, label in choices)
        else:
            slots.append((column, None))
    width = len(slots)

    _brand_row(main, f"ARK CRM  ·  {plural} import sample  ·  {generated:%d-%b-%Y}  ·  Fill in from row {FIRST_DATA_ROW}", width)

    # ---- row 2: one band per form section
    start = 1
    for index in range(1, width + 2):
        here = slots[index - 1][0].section if index <= width else None
        if index > width or here != slots[start - 1][0].section:
            cell = main.cell(row=2, column=start, value=section_title(slots[start - 1][0].section))
            for c in range(start, index):
                main.cell(row=2, column=c).fill = SECTION_FILL
            cell.font = SECTION_FONT
            if index - 1 > start:
                main.merge_cells(start_row=2, start_column=start, end_row=2, end_column=index - 1)
            start = index

    tallest_hint = tallest_condition = 1
    last_row = FIRST_DATA_ROW + TEMPLATE_ROWS - 1
    for index, (column, choice) in enumerate(slots, start=1):
        letter = get_column_letter(index)
        label = choice_header(column, choice) if choice else column.label
        must = column.api_name in required
        # A multi-choice field split into Yes/No columns is required as a GROUP
        # (at least one Yes), so its columns are shaded but carry no star — a
        # star on each read as "answer every one".
        star = must and not choice
        header = main.cell(row=LABEL_ROW, column=index, value=f"{label} *" if star else label)
        header.font = HEADER_FONT
        header.fill = REQUIRED_FILL if must else LATER_FILL if column.api_name in later else HEADER_FILL
        header.alignment = Alignment(vertical="center", wrap_text=True)
        col_width = max(14, min(30, len(label) + 3))
        main.column_dimensions[letter].width = col_width

        hint = how_to_fill(column, choice=bool(choice))
        cell = main.cell(row=HINT_ROW, column=index, value=hint)
        cell.font = HINT_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        tallest_hint = max(tallest_hint, _lines(hint, col_width * 1.15))

        condition = conditions.get(column.api_name)
        if condition:
            cell = main.cell(row=CONDITION_ROW, column=index, value=condition)
            cell.font = CONDITION_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            tallest_condition = max(tallest_condition, _lines(condition, col_width * 1.15))

        # ---- the cell check, with the hint shown when a cell is selected
        kind = column.type
        validation: DataValidation
        if choice or kind == "checkbox":
            validation = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
            validation.error = "Choose Yes or No."
        elif kind == "picklist" and (source := list_range(
            column.picklist or column.api_name, column.label,
            [option for _, option, active in directory.options(column.picklist) if active],
        )):
            validation = DataValidation(type="list", formula1=source, allow_blank=True)
            validation.error = "Pick one of the choices in the list."
        elif kind == "lookup" and (source := list_range(f"lookup:{column.api_name}", column.label, lookup_choices(column, directory))):
            validation = DataValidation(type="list", formula1=source, allow_blank=True, errorStyle="warning")
            validation.error = f"This name isn't in the list from {generated:%d-%b-%Y}. Keep it only if it was added to ARK since."
        elif kind in ("date", "datetime"):
            validation = DataValidation(type="date", operator="between", formula1=str(_EARLIEST_SERIAL), formula2=str(_LATEST_SERIAL), allow_blank=True)
            validation.error = "Enter a month like Oct-2026." if column.api_name.endswith("_month") else "Enter a date like 17-Sep-2026."
            number_format = MONTH_FORMAT if column.api_name.endswith("_month") else DATE_FORMAT
            for row in range(FIRST_DATA_ROW, last_row + 1):
                main.cell(row=row, column=index).number_format = number_format
        elif kind in ("number", "currency", "percent"):
            low = column.min_value if column.min_value is not None else (0 if kind == "percent" else -_NUMBER_BOUND)
            high = column.max_value if column.max_value is not None else (100 if kind == "percent" else _NUMBER_BOUND)
            validation = DataValidation(type="decimal", operator="between", formula1=f"{low:g}", formula2=f"{high:g}", allow_blank=True)
            validation.error = f"Enter a number between {low:,.0f} and {high:,.0f}, without symbols." if abs(high) < _NUMBER_BOUND else "Enter a number, without symbols."
            if kind == "currency":
                for row in range(FIRST_DATA_ROW, last_row + 1):
                    main.cell(row=row, column=index).number_format = MONEY_FORMAT
        elif kind == "text" and column.max_length:
            validation = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=str(column.max_length), allow_blank=True)
            validation.error = f"Keep it to {column.max_length} characters."
        else:
            validation = DataValidation(allow_blank=True)
        validation.showErrorMessage = validation.type is not None
        validation.errorTitle = label[:32]
        validation.showInputMessage = True
        validation.promptTitle = label[:32]
        validation.prompt = (hint + (f"\n{condition}" if condition else ""))[:255]
        main.add_data_validation(validation)
        validation.add(f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}")

    main.row_dimensions[LABEL_ROW].height = 30
    main.row_dimensions[HINT_ROW].height = 13 * tallest_hint + 2
    main.row_dimensions[CONDITION_ROW].height = 13 * tallest_condition + 2
    main.freeze_panes = f"B{FIRST_DATA_ROW}"

    # ---- Instructions
    _brand_row(instructions, f"ARK CRM  ·  How to import {plural.lower()}  ·  {generated:%d-%b-%Y}", 2)
    instructions.column_dimensions["A"].width = 16
    instructions.column_dimensions["B"].width = 100
    steps = [
        "Import in this order: Accounts (and Partners), then Contacts, then Leads. A row can only name accounts and people already in ARK.",
        f"Fill in the {main.title} sheet from row {FIRST_DATA_ROW}, one {noun} per row. Leave rows 1–5 as they are.",
        "Row 4 says what to enter in each column. Row 5 says when a column applies. Leave it blank otherwise.",
        "Dropdown columns: pick a value from the list.",
        "Columns named like “Account Type · End Client”: put Yes under every choice that applies.",
        f"“Pick from ARK” columns list the names in ARK on {generated:%d-%b-%Y}. A name added since then still imports.",
        "Dates: write 17-Sep-2026. A date like 03/04/2026 is refused, because it could mean two different days.",
        *notes,
        "ARK checks every row first and shows what needs fixing. Nothing is saved until every row is ready.",
    ]
    instructions.cell(row=3, column=1, value="How to use this file").font = Font(bold=True, size=12)
    row = 4
    for step in steps:
        instructions.cell(row=row, column=1, value="•").alignment = Alignment(horizontal="right", vertical="top")
        cell = instructions.cell(row=row, column=2, value=step)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    instructions.cell(row=row, column=1, value="Colours in row 3").font = Font(bold=True, size=12)
    row += 1
    for fill, name, meaning in (
        (REQUIRED_FILL, "Amber, with *", "Must be filled in, or the row isn't imported."),
        (LATER_FILL, "Green", "Needed before the record moves to its next stage in ARK. Fill it in now if you know it."),
        (HEADER_FILL, "Grey", "Optional."),
    ):
        swatch = instructions.cell(row=row, column=1, value=name)
        swatch.fill = fill
        swatch.font = HEADER_FONT
        instructions.cell(row=row, column=2, value=meaning)
        row += 1

    if list_ranges:
        lists.sheet_state = "hidden"
    else:
        book.remove(lists)
    book.active = book.sheetnames.index(main.title)
    instructions.sheet_view.tabSelected = False
    main.sheet_view.tabSelected = True
    return save(book)


def template_csv(columns: list[Column], required: set[str]) -> bytes:
    """The CSV sample: one header row, for another system's export to line up with."""
    text = io.StringIO()
    csv.writer(text).writerow([f"{c.label} *" if c.api_name in required else c.label for c in columns])
    return ("﻿" + text.getvalue()).encode("utf-8")


def write_table(sheet, headers: list[str], rows: Iterable[list[Any]], kinds: list[str], *, banner: str) -> None:
    """An export sheet: the Astrikos band on row 1, labels on row 2, then the rows."""
    _brand_row(sheet, banner, len(headers))
    sheet.append(headers)
    for index in range(1, len(headers) + 1):
        cell = sheet.cell(row=2, column=index)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "C3" if len(headers) > 2 else "A3"
    for row in rows:
        sheet.append([_excel_value(v) for v in row])
        excel_row = sheet.max_row
        for index, kind in enumerate(kinds, start=1):
            cell = sheet.cell(row=excel_row, column=index)
            _as_text(cell)
            if kind in ("date", "datetime") and isinstance(cell.value, (date, datetime)):
                cell.number_format = DATE_FORMAT
            elif kind == "currency" and isinstance(cell.value, (int, float)):
                cell.number_format = MONEY_FORMAT
    _autosize(sheet)


def _autosize(sheet, minimum: int = 10, maximum: int = 48) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows(min_row=2, max_row=min(sheet.max_row, 200)):
        for cell in row:
            if cell.value is None:
                continue
            length = len(cell.value.strftime("%d-%b-%Y")) if isinstance(cell.value, (date, datetime)) else len(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), min(length, maximum))
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = max(minimum, width + 2)


def _excel_value(value: Any) -> Any:
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def new_book() -> Workbook:
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = "Astrikos · ARK CRM"
    return book


def save(book: Workbook) -> bytes:
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
